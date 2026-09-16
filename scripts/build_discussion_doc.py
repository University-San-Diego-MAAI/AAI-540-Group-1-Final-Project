"""Build the Canvas discussion RFC proposal post (discussion.docx).

Word-count-sensitive deliverable: the course limit is < 1,000 words. The script
counts every word in the generated document (title + headings + body) and fails
loudly above the cap.

Run from the repo root:  .venv/bin/python scripts/build_discussion_doc.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "discussion.docx"
WORD_CAP = 1000

ACCENT = RGBColor(0x1F, 0x4E, 0x79)

TITLE = (
    "RFC Proposal: Predicting Health Insurance Member Utilization Risk "
    "Using Multiclass Classification (AAI-540, Group 1)"
)

SECTIONS = [
    ("Project Summary", [
        "Healthcare insurers need to know which members are likely to require "
        "different levels of healthcare utilization and care-management attention "
        "in an upcoming period. Catching a high-risk member early - before "
        "avoidable admissions accumulate - is far cheaper than reacting after "
        "the fact. Our team proposes an ML system, for a hypothetical "
        "UnitedHealthcare scenario, that predicts each member's utilization-risk "
        "category from historical claims, prescription, and utilization data so "
        "that care-management resources (nurse outreach, care coordinators, "
        "disease-management programs) can be directed proactively. The business "
        "impact is concrete: every High-Risk member identified early is an "
        "opportunity for intervention before high-cost utilization occurs, while "
        "over-flagging Low-Risk members only costs outreach effort.",
    ]),
    ("What is the model's objective?", [
        "The objective is a supervised multiclass classifier that assigns every "
        "member a Member_Utilization_Risk label - Low, Medium, or High - for an "
        "upcoming period, given that member's history-window profile: "
        "demographics, chronic conditions, inpatient/outpatient utilization, "
        "physician/carrier services, prescription utilization, and insurance "
        "coverage. Because claims data arrives with no risk labels, the target "
        "is constructed with a transparent rule: features come from a 2008 "
        "history window, and the label is derived from observed 2009-2010 "
        "outcome utilization (a weighted score of total visits, inpatient "
        "admissions, and drug eras, terciled into the three risk tiers). The "
        "83.3% of members observable in both windows form the labeled set.",
    ]),
    ("What type of Machine Learning problem will you be solving?", [
        "A supervised, multiclass (three-class) classification problem on tabular "
        "member-level data. We aggregate event-level claims records into one row "
        "per member (24 features: age band, gender, race, state; visit counts by "
        "type, care sites, observation months; six chronic-condition flags from "
        "ICD-9 chapters; drug exposure/era volumes; distinct conditions and "
        "procedures). Candidate models are multinomial logistic regression "
        "(interpretable baseline), Random Forest, and XGBoost (primary), selected "
        "on validation performance, interpretability, and suitability for the "
        "business objective.",
    ]),
    ("What is your data source?", [
        "The CMS Data Entrepreneurs' Synthetic Public Use File (DE-SynPUF), "
        "2008-2010, converted to the OHDSI OMOP Common Data Model v5.x and hosted "
        "as a public AWS Open Data dataset (AWS Marketplace listing: "
        "https://aws.amazon.com/marketplace/pp/prodview-f7br22sw3qlg2; S3 bucket "
        "s3://synpuf-omop, us-east-1; no credentials or subscription required). "
        "We use the 100,000-person sample across 17 OMOP tables (~915 MB "
        "compressed): person (100k rows), visit_occurrence (4.79M), "
        "condition_occurrence (12.70M), drug_exposure (5.42M), procedure_"
        "occurrence (11.90M), observation_period, payer_plan_period, and death, "
        "all linkable by person_id. We selected it because it is the only public "
        "dataset with the true structure of CMS Medicare claims - beneficiary, "
        "inpatient/outpatient, carrier, and prescription domains - without the "
        "cost, privacy, and Data Use Agreement barriers of real CMS data.",
        "Key caveats we have already quantified in exploratory analysis: the data "
        "is fully synthetic (CMS states it has limited inferential research "
        "value); the OMOP conversion contains no payment columns, so 'previous "
        "spending' is proxied by utilization intensity; 85% of visit rows have "
        "unmapped visit-type concepts; and 10.7% of members have no observation "
        "period, so raw zero-utilization partly means 'not observed' rather than "
        "'low risk' - we control for this with exposure-time features.",
    ]),
    ("How will you evaluate your model?", [
        "With complementary multiclass metrics rather than accuracy alone. "
        "Primary metric: macro F1-score, so all three risk tiers count equally. "
        "Supporting metrics: per-class precision and recall (recall on the "
        "High-Risk class is the headline business number - missing a High-Risk "
        "member is more consequential than over-flagging), the confusion matrix "
        "to inspect which tiers are confused, and multiclass ROC-AUC "
        "(one-vs-rest). Data is split with stratified 70/15/15 "
        "train/validation/test partitions plus k-fold cross-validation during "
        "hyperparameter tuning, and we additionally audit fairness slices "
        "(gender, race, age band) so performance gaps across member groups are "
        "visible before deployment.",
    ]),
    ("How will you monitor your model?", [
        "On three axes, all feeding CloudWatch dashboards and alarms. (1) Model: "
        "prediction-drift tracking - the share of members scored Low/Medium/High "
        "each run is compared against the training baseline with PSI-style "
        "alerts - and, once outcome windows mature, recomputed ground-truth "
        "macro F1 and per-class recall over time to trigger retraining on decay. "
        "(2) Data: an automated validation gate on every new extract (schema "
        "checks, row counts, primary-key uniqueness, person_id referential "
        "integrity), null-rate thresholds, feature-distribution drift versus the "
        "training snapshot, and explicit tracking of observation-window coverage, "
        "since the not-observed population is the main data-quality trap in this "
        "dataset. (3) Infrastructure: SageMaker job success/failure, processing "
        "and batch-scoring durations, S3 metrics, and pipeline execution status, "
        "with SNS notifications to the team on any failure.",
    ]),
    ("Why this problem is a good fit - and what makes it non-trivial", [
        "The dataset and the problem statement align closely: the RFC calls for "
        "beneficiary, inpatient/outpatient, carrier, and prescription domains "
        "linkable by a single member identifier, and the OMOP conversion provides "
        "exactly those domains keyed by person_id, which our validation gate "
        "confirms is orphan-free. The task is non-trivial rather than synthetic-"
        "looking: utilization counts are heavily right-skewed with a zero-inflated "
        "tail, chronic conditions correlate with but are not determined by raw "
        "visit volume, and our tercile-proxy preview shows the three risk groups "
        "overlap on individual features while separating cleanly in aggregate - "
        "so model selection and feature engineering will matter. The main "
        "challenges we anticipate are label sensitivity (the tercile weighting "
        "shapes what 'High Risk' means), exposure-time confounding, and the "
        "absence of cost fields; each is addressed in the design document with a "
        "documented mitigation rather than ignored.",
    ]),
    ("Next Steps", [
        "Our pipeline (data validation gate, feature engineering, temporal "
        "splitting) is prototyped locally; we are now porting it to SageMaker "
        "(Processing, Training, Batch Transform) with the CI/CD pipeline, "
        "evaluation gates, and monitoring from our ML design document. Project "
        "repository: https://github.com/University-San-Diego-MAAI/AAI-540-"
        "Group-1-Final-Project. We welcome feedback on the labeling rule and the "
        "risk-tier weighting.",
    ]),
]


def set_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size in [("Heading 1", 14), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.font.bold = True


def count_words(doc: Document) -> int:
    words = 0
    for para in doc.paragraphs:
        words += len(para.text.split())
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                words += len(cell.text.split())
    return words


def build() -> None:
    doc = Document()
    set_styles(doc)

    title = doc.add_paragraph()
    run = title.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(15)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for heading, paragraphs in SECTIONS:
        doc.add_heading(heading, level=2)
        for text in paragraphs:
            doc.add_paragraph(text)

    total = count_words(doc)
    print(f"Word count: {total} (cap {WORD_CAP})")
    if total >= WORD_CAP:
        raise SystemExit(f"FAIL: {total} words exceeds the {WORD_CAP} cap")

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
